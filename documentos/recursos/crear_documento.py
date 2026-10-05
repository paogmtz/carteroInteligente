from pathlib import Path
import sys, itertools
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.path.insert(0,str(Path.cwd()))
import main
ROOT=Path.cwd()/'documentos'; AS=ROOT/'recursos'
ruta,costo=main.construir_ruta()
assert len(ruta)==len(set(ruta))==20 and costo==35
lines=Path('main.py').read_text().splitlines()
fontpath='/System/Library/Fonts/Menlo.ttc'
font=ImageFont.truetype(fontpath,26)
for idx,(a,b) in enumerate([(1,51),(54,76),(79,112)],1):
    rows=[(n,lines[n-1]) for n in range(a,b+1) if lines[n-1].strip()]
    im=Image.new('RGB',(1510,58+37*len(rows)), '#f5f7fa'); dr=ImageDraw.Draw(im)
    for j,(n,line) in enumerate(rows):
        y=25+j*37
        dr.text((18,y),f'{n:3}',font=font,fill='#7b8794')
        color='#23714c' if line.lstrip().startswith(('#','"""')) else '#182537'
        dr.text((98,y),line,font=font,fill=color)
    im.save(AS/f'codigo_{idx}.png')
colors=['#eab75b','#8cb9dd','#a9cb99','#d5a1b3','#b9aae1']
zone_of={n:i for i,z in enumerate(main.zonas) for n in z}
def af(size): return ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',size)
def center(dr,xy,t,size=25,fill='#182537'):
    dr.text(xy,t,font=af(size),fill=fill,anchor='mm')
def arrow(dr,a,b,color='#165482',width=5,shrink=0):
    dx,dy=b[0]-a[0],b[1]-a[1]; L=math.hypot(dx,dy); ux,uy=dx/L,dy/L
    a=(a[0]+ux*shrink,a[1]+uy*shrink); b=(b[0]-ux*shrink,b[1]-uy*shrink)
    dr.line([a,b],fill=color,width=width)
    dr.polygon([b,(b[0]-ux*17-uy*8,b[1]-uy*17+ux*8),(b[0]-ux*17+uy*8,b[1]-uy*17-ux*8)],fill=color)
im=Image.new('RGB',(1600,1110),'white'); dr=ImageDraw.Draw(im)
coords={n:(100+c*138,90+f*138) for n,(f,c) in main.posiciones.items()}
for a,b in itertools.combinations(main.posiciones,2): dr.line([coords[a],coords[b]],fill='#eef0f3',width=1)
for a,b in zip(ruta,ruta[1:]):
    A,B=coords[a],coords[b]
    x=(A[0]+B[0])/2; y=(A[1]+B[1])/2-22
    if (a,b)==(8,12):
        bend=(x,A[1]-80)
        dr.line([(A[0]-31,A[1]-9),bend],fill='#165482',width=5)
        arrow(dr,bend,B,shrink=34)
        y=A[1]-100
    else: arrow(dr,A,B,shrink=34)
    dr.rectangle((x-15,y-15,x+15,y+15),fill='white');center(dr,(x,y),str(main.distancia(a,b)),28,'#165482')
for n,(x,y) in coords.items():
    dr.ellipse((x-31,y-31,x+31,y+31),fill=colors[zone_of[n]],outline='#334155',width=2);center(dr,(x,y),str(n),28)
for i in range(5):
    x=240+i*255; dr.ellipse((x-15,1035,x+15,1065),fill=colors[i]); dr.text((x+25,1036),'Zona '+chr(65+i),font=af(28),fill='black')
center(dr,(800,985),'Posiciones relativas de los destinos en la matriz',27)
im.save(AS/'grafo.png')
im=Image.new('RGB',(1450,1530),'white');dr=ImageDraw.Draw(im);acc=0
for k,n in enumerate(ruta):
    y=55+k*74
    if k:
        step=main.distancia(ruta[k-1],n);acc+=step;arrow(dr,(390,y-51),(390,y-25),width=3)
        center(dr,(270,y-37),'+'+str(step),22,'#165482')
    state=f'({n}, V{k})' if k else '(10, {10})'
    dr.rounded_rectangle((285,y-23,495,y+23),radius=9,fill='#e7f0f7',outline='#165482',width=2)
    center(dr,(390,y),state,26);dr.text((520,y-14),f'g = {acc}',font=af(24),fill='#334155')
for x,n in [(860,11),(1240,19)]:
    arrow(dr,(495,60),(x,106),color='#999999',width=2)
    dr.rounded_rectangle((x-160,106,x+160,154),radius=9,fill='#f5f5f5',outline='#999999',width=2)
    center(dr,(x,130),f'({n}, {{10, {n}}})',25);center(dr,(x,205),'...',35,'#777777')
center(dr,(1050,305),'Alternativas iniciales no ejecutadas',27,'#666666')
for i,t in enumerate(['Vk = conjunto de los primeros','k + 1 nodos de la ruta.','','Ejemplo','V3 = {10, 9, 11, 19}','','g = costo acumulado.']):center(dr,(1040,540+i*45),t,28)
for i,t in enumerate(['Meta','V19 contiene los 20 nodos','y el costo final es 35.']):center(dr,(1040,1140+i*45),t,28)
im.save(AS/'arbol.png')
D=Document(); sec=D.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11); sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.75)
for st in ['Normal','Title','Heading 1','Heading 2']:
    D.styles[st].font.name='Calibri'; D.styles[st].font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(11); D.styles['Normal'].paragraph_format.space_after=Pt(7)
D.styles['Normal'].paragraph_format.line_spacing=1.08
D.styles['Title'].font.size=Pt(25); D.styles['Heading 1'].font.size=Pt(17)
D.styles['Heading 2'].font.size=Pt(12)
def p(t): return D.add_paragraph(t)
def h(t): D.add_heading(t,1)
def img(name,width=7): D.add_picture(str(AS/name),width=Inches(width))
def page(): D.add_page_break()
def lead(label,text):
    a=D.add_paragraph(); a.add_run(label+' ').bold=True; a.add_run(text)
D.add_heading('El cartero inteligente',0)
p('Explicación de nuestra propuesta y del código para el equipo')
p('Buscamos una ruta que salga del Palacio de Correos, nodo 10, y visite los 20 destinos del Centro Histórico de la Ciudad de México con el menor costo posible. En esta primera versión no pedimos regresar al inicio. Nuestra propuesta usa una matriz y zonas para obtener una ruta sencilla de calcular; el resultado actual cuesta 35 unidades, sin garantía de ser la mejor ruta global.')
h('Cómo formulamos el problema')
table=D.add_table(rows=1,cols=2); table.autofit=False; table.columns[0].width=Inches(1.45); table.columns[1].width=Inches(5.55)
table.rows[0].cells[0].text='Componente'; table.rows[0].cells[1].text='Definición del equipo'
records=[('Estado inicial','(10, {10}). El cartero empieza en el Palacio de Correos y el nodo 10 ya cuenta como visitado.'),('Estados','(actual, visitados): destino actual y conjunto de destinos ya atendidos. Estar en el mismo nodo con diferentes visitados son estados distintos.'),('Acciones','Moverse a un destino pendiente y añadirlo a visitados. La estrategia actual termina primero los pendientes de la zona donde se encuentra.'),('Estado meta','Haber visitado los 20 nodos. Puede terminar en cualquier destino; la ruta actual acaba en 18.'),('Función de costo','Sumar las distancias Manhattan de cada salto: |fila1 − fila2| + |columna1 − columna2|. Son unidades relativas de la matriz.'),('Restricciones','Inicio en 10; todos los destinos deben visitarse. El código elige cada destino una vez. Los ceros son espacios vacíos; no modelamos calles, obstáculos, tráfico ni sentidos de circulación.')]
for a,b in records:
    cells=table.add_row().cells; cells[0].text=a; cells[1].text=b
for i,row in enumerate(table.rows):
    for c in row.cells:
        pr=c._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'DCE6F1' if i==0 else ('F5F7FA' if i%2 else 'FFFFFF')); pr.append(shade)
        borders=OxmlElement('w:tcBorders')
        for side in ['top','left','bottom','right']:
            el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
        pr.append(borders)
        margins=OxmlElement('w:tcMar');
        for side in ['top','left','bottom','right']:
            el=OxmlElement('w:'+side); el.set(qn('w:w'),'100'); el.set(qn('w:type'),'dxa'); margins.append(el)
        pr.append(margins)
        c.vertical_alignment=1
        for pp in c.paragraphs:
            pp.paragraph_format.space_after=Pt(3)
            for r in pp.runs: r.font.size=Pt(10.5); r.bold=i==0
h('Por qué usamos zonas')
p('Agrupamos destinos cercanos para organizar las decisiones y evitar saltos constantes entre sectores. A = {12, 13, 14, 15}; B = {2, 3, 4, 5, 20}; C = {1, 6, 7, 8}; D = {9, 10, 11, 19}; E = {16, 17, 18}. Son conjuntos dentro de una lista, sin grafos anidados. Esta versión completa cada zona antes de salir; para futuros algoritmos esa regla puede quitarse.')
page(); h('La matriz y las posiciones')
p('La matriz conserva de forma aproximada la distribución del mapa. Cada número identifica un destino y cada 0 deja un espacio vacío. No es un mapa de calles: la distancia entre destinos se calcula directamente con sus coordenadas.')
img('codigo_1.png',6.65)
p('Captura del código actual, líneas 1 a 51. Se omiten únicamente las líneas vacías; la numeración corresponde a main.py. Las tres capturas incluyen todo el código no vacío.')
lead('mapa y zonas', 'mapa guarda filas y columnas; zonas agrupa números de destinos. permutations se importa para comparar órdenes de visita más adelante.')
lead('posiciones', 'Los dos for recorren la matriz. Si el valor no es 0, guardan nodo: (fila, columna). Los índices empiezan en 0; por ejemplo, posiciones[10] = (2, 5).')
lead('distancia y obtener_zona', 'distancia suma los cambios horizontal y vertical. De 10 a 9 cuesta |2 − 1| + |5 − 5| = 1. obtener_zona recorre la lista y devuelve el conjunto al que pertenece el nodo.')
p('Las posiciones se conservan tal como están en el código: el 7 queda arriba y a la derecha del 8. Ajustar la matriz puede cambiar el costo y la ruta.')
page(); h('Cómo elegimos el recorrido de una zona')
p('La idea original era reservar un nodo cercano al exterior como salida. La versión actual compara también lo que cuesta recorrer la zona antes de salir. Así evitamos elegir una salida conveniente que obligue a dar una vuelta interna costosa.')
img('codigo_2.png')
lead('costo_recorrido', 'Forma una secuencia con el inicio y el orden propuesto. zip empareja cada nodo con el siguiente y sum suma los costos de esos saltos.')
lead('pendientes y fuera', 'La resta de conjuntos quita los visitados. pendientes contiene los destinos que faltan en la zona; fuera contiene los que faltan en otras zonas.')
lead('puntuacion', 'Para cada orden suma el costo interno y el salto desde su último nodo al pendiente exterior más cercano. orden[-1] es el último nodo. Si ya no queda nadie fuera, default=0 evita sumar una salida innecesaria.')
lead('permutations y min', 'permutations genera los órdenes posibles; min elige el de menor puntuación. En empates gana el menor costo interno y luego el orden numérico. sorted ayuda a obtener resultados reproducibles.')
p('Ejemplo desde 10: el orden 9 → 11 → 19 cuesta 1 + 3 + 2 = 6 dentro de la zona D. Salir de 19 hacia 4 añade 2. Su puntuación es 8. El salto exterior se usa para evaluar la decisión; el costo real se suma cuando el cartero se mueve.')
p('Con cinco pendientes hay como máximo 5! = 120 órdenes. Es una comparación pequeña por zona, no una evaluación de todas las rutas de los 20 destinos. La decisión sigue siendo local.')
page(); h('Cómo construimos y mostramos la ruta')
img('codigo_3.png')
lead('Estado y registro', 'actual es el destino donde estamos; visitados es un conjunto sin duplicados. Juntos representan el estado conceptual. ruta guarda el orden y costo guarda la distancia acumulada.')
lead('while e if', 'El ciclo continúa hasta visitar todos los destinos. Si quedan pendientes en la zona, toma el orden completo elegido. Si la zona terminó, elige el pendiente más cercano; los empates se resuelven por número de nodo.')
lead('for y return', 'Para cada movimiento suma su distancia, cambia actual, añade el nodo a visitados y lo registra en ruta. Al final devuelve ruta y costo. El bloque __main__ imprime el resultado solo al ejecutar el archivo directamente.')
p('Ruta obtenida: 10 → 9 → 11 → 19 → 4 → 20 → 5 → 2 → 3 → 1 → 6 → 7 → 8 → 12 → 13 → 14 → 15 → 16 → 17 → 18. Costo total: 35 unidades.')
p('Se verificó que empieza en 10, incluye los 20 destinos una sola vez y que la suma de sus 19 movimientos es 35. No se calcula un regreso de 18 a 10.')
page(); h('El grafo del modelo y la ruta resultante')
p('Cada destino es un nodo. Como distancia(a, b) se puede calcular para cualquier par, el modelo representa un grafo completo, no dirigido y con pesos: 20 nodos y 190 aristas. No hace falta almacenar esas aristas; sus pesos se obtienen de la matriz.')
img('grafo.png')
p('Las líneas grises muestran las conexiones posibles. Las flechas azules resaltan los 19 movimientos de la ruta elegida y sus números indican el costo Manhattan. Los colores identifican las zonas. Una línea diagonal solo conecta dos destinos: no representa una calle diagonal ni un trayecto real.')
lead('Cómo leerlo', 'La primera flecha va de 10 a 9 y cuesta 1. Después, 9 a 11 cuesta 3. El recorrido pasa por las zonas D → B → C → A → E y termina en 18.')
p('Costos de los movimientos en orden: 1, 3, 2, 2, 2, 2, 3, 1, 2, 2, 1, 2, 4, 1, 1, 1, 3, 1, 1. Su suma es 35.')
p('La propuesta simplifica el mapa original: mide cercanía horizontal y vertical, pero no verifica por cuáles calles se puede circular. Para usar distancias reales habría que agregar esa información al modelo.')
page(); h('El árbol de estados y la rama elegida')
p('En un árbol, cada nodo representa (actual, visitados), y cada arista representa una decisión de movimiento. Aquí mostramos completa la rama ejecutada y dos alternativas iniciales de la zona D. Las demás ramas se omiten para que el dibujo sea legible.')
img('arbol.png',6.15)
p('Este es un árbol conceptual parcial. El código compara órdenes locales y registra una ruta; todavía no construye ni explora un árbol con BFS, DFS, Costo Uniforme o A*. El grafo muestra destinos; el árbol distingue el historial de visitados.')
p('Después podemos reutilizar matriz, posiciones y distancia. BFS minimiza transiciones; Costo Uniforme minimiza costo; A* necesita una heurística revisada. Para comparar rutas globales conviene permitir salir de una zona aun cuando queden destinos en ella.')
for element in D.styles.element.iter():
    for child in list(element):
        if child.tag == qn('w:pBdr'): element.remove(child)
D.core_properties.title='El cartero inteligente'; D.core_properties.subject='Propuesta y explicación del código'; D.core_properties.author='Equipo del proyecto'
D.save(ROOT/'El_cartero_inteligente.docx')
print(ROOT/'El_cartero_inteligente.docx')
